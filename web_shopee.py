# ==============================================================================
# BCM CLOUD v4.9.1 - STABLE VERSION (FIX BUGS)
# Coder: BCM-Engineer (An) & Sếp Lâm
# Update:
# 1. Sửa lỗi đỏ (UnboundLocalError) khi chưa upload file.
# 2. Tích hợp đầy đủ: Word Export, Excel Import, Backup.
# ==============================================================================

import streamlit as st
import sqlite3
import pandas as pd
from datetime import datetime, timedelta
import os
import google.generativeai as genai
from pypdf import PdfReader
from docx import Document
import re
import io

# ==================================================
# 1. CẤU HÌNH HỆ THỐNG
# ==================================================
st.set_page_config(page_title="BCM Cloud v4.9.1 - MIT Corp", page_icon="🦅", layout="wide")
st.markdown("""<style>.stMetric {background-color: #f0f2f6; padding: 10px; border-radius: 5px;} [data-testid="stMetricValue"] {font-size: 1.5rem !important;}</style>""", unsafe_allow_html=True)

# Lấy API Key
AI_STATUS = "Offline 🔴"
try:
    if "GOOGLE_API_KEY" in st.secrets:
        genai.configure(api_key=st.secrets["GOOGLE_API_KEY"])
        AI_STATUS = "Online 🟢"
except: pass

MODEL_NAME = "gemini-3.1-pro-preview"
DB_FILE = "shopee_data_v3.db"

# ==================================================
# 2. HÀM DATABASE
# ==================================================
def init_db():
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS products (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT, cost_price INTEGER, selling_price INTEGER, stock_quantity INTEGER DEFAULT 0, alert_threshold INTEGER DEFAULT 5, daily_sales REAL DEFAULT 1.0, lead_time INTEGER DEFAULT 15, safety_stock INTEGER DEFAULT 5)''')
    c.execute('''CREATE TABLE IF NOT EXISTS financials (date TEXT PRIMARY KEY, revenue INTEGER DEFAULT 0, ad_spend INTEGER DEFAULT 0, profit INTEGER DEFAULT 0)''')
    c.execute('''CREATE TABLE IF NOT EXISTS competitors (comp_id INTEGER PRIMARY KEY AUTOINCREMENT, my_product_name TEXT, comp_name TEXT, comp_url TEXT, comp_price INTEGER, last_check TEXT)''')
    conn.commit(); conn.close()
init_db()

def get_products_df(): conn=sqlite3.connect(DB_FILE); df=pd.read_sql_query("SELECT * FROM products", conn); conn.close(); return df
def get_products_list(): df=get_products_df(); return df['name'].tolist() if not df.empty else []

def add_product_full(name, cost, price, stock, daily, lead, safe): 
    threshold = int(daily * lead + safe)
    conn = sqlite3.connect(DB_FILE)
    cur = conn.cursor()
    cur.execute("SELECT id FROM products WHERE name = ?", (name,))
    exists = cur.fetchone()
    if not exists:
        cur.execute("INSERT INTO products (name, cost_price, selling_price, stock_quantity, daily_sales, lead_time, safety_stock, alert_threshold) VALUES (?,?,?,?,?,?,?,?)", 
                    (name, cost, price, stock, daily, lead, safe, threshold))
    else:
        cur.execute("""
            UPDATE products 
            SET cost_price=?, selling_price=?, stock_quantity=?, daily_sales=?, lead_time=?, safety_stock=?, alert_threshold=? 
            WHERE name=?""", 
            (cost, price, stock, daily, lead, safe, threshold, name))
    conn.commit(); conn.close()

def add_competitor(m,c,u,p): d=datetime.now().strftime("%Y-%m-%d"); conn=sqlite3.connect(DB_FILE); cur=conn.cursor(); cur.execute("INSERT INTO competitors (my_product_name,comp_name,comp_url,comp_price,last_check) VALUES (?,?,?,?,?)",(m,c,u,p,d)); conn.commit(); conn.close()
def get_competitors_df(): conn=sqlite3.connect(DB_FILE); df=pd.read_sql_query("SELECT * FROM competitors", conn); conn.close(); return df

def get_file_content(uploaded_file):
    text = ""
    try:
        if uploaded_file.name.endswith(".pdf"):
            pdf_reader = PdfReader(uploaded_file); 
            for page in pdf_reader.pages: text += page.extract_text() + "\n"
        elif uploaded_file.name.endswith(".docx"):
            doc = Document(uploaded_file); 
            for para in doc.paragraphs: text += para.text + "\n"
        elif uploaded_file.name.endswith(".txt"): text = uploaded_file.read().decode("utf-8")
    except: pass
    return text

# ==================================================
# 3. HÀM TẠO FILE WORD
# ==================================================
def create_word_docx(role_name, content):
    doc = Document()
    doc.add_heading('BIÊN BẢN HỌP CHIẾN LƯỢC - BCM CLOUD', 0)
    p = doc.add_paragraph()
    p.add_run(f'Thời gian: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}\n').bold = True
    p.add_run(f'Tham vấn: {role_name}\n').bold = True
    p.add_run('--------------------------------------------------')
    doc.add_heading('NỘI DUNG PHÂN TÍCH:', level=1)
    for line in content.split('\n'):
        doc.add_paragraph(line)
    doc.add_paragraph('--------------------------------------------------')
    doc.add_paragraph('Báo cáo được tạo tự động bởi hệ thống BCM Cloud.')
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ==================================================
# 4. LOGIC XỬ LÝ SỐ LIỆU SHOPEE
# ==================================================
def parse_vn_currency(val):
    if pd.isna(val): return 0
    s = str(val).strip()
    s = re.sub(r'[^\d.,]', '', s) 
    if '.' in s and ',' in s: s = s.replace('.', '').replace(',', '.')
    elif '.' in s:
        parts = s.split('.')
        if len(parts) > 1 and (len(parts) > 2 or len(parts[-1]) == 3): s = s.replace('.', '')
    elif ',' in s: s = s.replace(',', '.')
    try: return float(s)
    except: return 0.0

def find_best_column(columns, keywords, blacklist=[]):
    cols_lower = [str(c).lower().strip() for c in columns]
    for kw in keywords:
        if kw in cols_lower: return columns[cols_lower.index(kw)]
    for col in columns:
        c_low = str(col).lower()
        if not any(k in c_low for k in keywords): continue
        if any(b in c_low for b in blacklist): continue
        return col
    return None

def process_shopee_files(revenue_file, ads_file):
    total_rev = 0; total_ads = 0; logs = []
    if revenue_file:
        try:
            revenue_file.seek(0)
            if revenue_file.name.endswith(('xls', 'xlsx')): df = pd.read_excel(revenue_file, header=0, dtype=str)
            else: df = pd.read_csv(revenue_file, header=0, dtype=str, encoding='utf-8')
        except: logs.append("❌ Lỗi đọc file Doanh thu"); df = pd.DataFrame()

        if not df.empty:
            col_rev = find_best_column(df.columns, keywords=["tổng doanh số (vnd)", "doanh số (vnd)", "tổng tiền", "doanh thu"], blacklist=["thẻ sản phẩm", "livestream", "video"])
            if col_rev:
                val = df[col_rev].iloc[0]
                total_rev = parse_vn_currency(val)
                logs.append(f"✅ Doanh thu: {total_rev:,.0f}")
            else: logs.append(f"⚠️ Không tìm thấy cột Doanh thu.")

    if ads_file:
        try:
            ads_file.seek(0)
            if ads_file.name.endswith(('xls', 'xlsx')): df_ads = pd.read_excel(ads_file, skiprows=6, dtype=str)
            else:
                try: df_ads = pd.read_csv(ads_file, skiprows=6, dtype=str, encoding='utf-8')
                except: df_ads = pd.read_csv(ads_file, skiprows=6, dtype=str, encoding='utf-16', sep='\t')
        except: logs.append("❌ Lỗi đọc file Ads"); df_ads = pd.DataFrame()

        if not df_ads.empty:
            col_cost = find_best_column(df_ads.columns, keywords=["chi phí", "cost"], blacklist=["chuyển đổi", "trực tiếp", "mỗi lượt", "roas"])
            if col_cost:
                total_ads = df_ads[col_cost].apply(parse_vn_currency).sum()
                logs.append(f"✅ Ads: {total_ads:,.0f}")
            else: logs.append(f"⚠️ Không tìm thấy cột Chi phí.")

    return total_rev, total_ads, logs

# ==================================================
# 5. GIAO DIỆN CHÍNH
# ==================================================
with st.sidebar:
    st.title("🦅 BCM Cloud v4.9.1")
    st.caption(f"Engine: {MODEL_NAME} | Status: {AI_STATUS}")
    st.markdown("---")
    menu = st.radio("Menu:", ["🤖 Phòng Họp Chiến Lược", "📊 Báo Cáo & Excel", "⚔️ Rada Đối Thủ", "💰 Tính Lãi & Nhập Kho", "📦 Kho Hàng & Backup"])
    
    if menu == "🤖 Phòng Họp Chiến Lược":
        st.markdown("---")
        st.subheader("📂 RAG (Nạp tài liệu)")
        uploaded_files = st.file_uploader("Upload PDF/Word:", accept_multiple_files=True, type=['pdf', 'docx', 'txt'])
        knowledge_context = ""
        if uploaded_files:
            with st.status("Đang học...", expanded=True) as status:
                for file in uploaded_files:
                    c = get_file_content(file)
                    if c: knowledge_context += f"\n--- DOC: {file.name} ---\n{c}\n"
                status.update(label="Đã học xong!", state="complete", expanded=False)

# ==================================================
# 6. LOGIC MODULES
# ==================================================
if menu == "📊 Báo Cáo & Excel":
    st.title("📊 BÁO CÁO KINH DOANH")
    d = st.date_input("Chọn tuần:", datetime.now())
    with st.expander("📂 UPLOAD FILE SHOPEE", expanded=True):
        f1 = st.file_uploader("File Doanh Thu (Shop Stats)")
        f2 = st.file_uploader("File Quảng Cáo (Ads)")
        
        # --- KHỞI TẠO BIẾN ĐỂ TRÁNH LỖI KHI CHƯA CÓ FILE ---
        rev = 0
        ads = 0
        debug_info = []
        
        if f1 or f2:
            rev, ads, debug_info = process_shopee_files(f1, f2)
            with st.expander("🔍 Log Xử Lý"):
                for l in debug_info: st.write(l)
                
    st.divider()
    c1, c2, c3 = st.columns(3)
    nr = c1.number_input("Doanh thu", float(rev), step=1e5, format="%.0f")
    na = c2.number_input("Chi phí Ads", float(ads), step=5e4, format="%.0f")
    np = c3.number_input("Lợi nhuận Ròng (30%)", float(nr*0.3-na), step=5e4, format="%.0f")
    
    # Export Excel
    output = io.BytesIO()
    data = {'Ngày': [datetime.now().strftime("%Y-%m-%d")], 'Doanh Thu': [nr], 'Ads': [na], 'Lợi Nhuận': [np]}
    df_export = pd.DataFrame(data)
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_export.to_excel(writer, index=False, sheet_name='BaoCao')
    excel_data = output.getvalue()
    st.download_button("💾 TẢI BÁO CÁO (.xlsx)", excel_data, "bao_cao_ngay.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")

elif menu == "🤖 Phòng Họp Chiến Lược":
    st.header("🤖 PHÒNG HỌP CHIẾN LƯỢC")
    df_comp = get_competitors_df()
    comp_context = f"\n--- THỊ TRƯỜNG ---\n{df_comp.to_string()}\n" if not df_comp.empty else ""
    
    col_role, col_info = st.columns([1, 3])
    with col_role: role = st.radio("Active:", ["An (Kỹ sư)", "Sư (Cố vấn)"], label_visibility="collapsed")
    with col_info:
        if "An" in role:
            st.info("**🔵 AN (ENGINEER):** Giải pháp kỹ thuật, tính toán, code.")
            prefix = "[🤖 Kỹ Sư AN]:"
            style = "Bạn là An. Trả lời ngắn gọn, kỹ thuật, con số."
            current_role_name = "Kỹ Sư An"
        else:
            st.warning("**🟠 SƯ (ADVISOR):** Chiến lược, phản biện, rủi ro.")
            prefix = "[👺 Quân Sư]:"
            style = "Bạn là Quân Sư. Soi xét, tìm rủi ro, chiến lược."
            current_role_name = "Quân Sư"

    st.divider()
    if "messages" not in st.session_state: st.session_state.messages = []
    for msg in st.session_state.messages: st.chat_message(msg["role"]).markdown(msg["content"])
    
    if p := st.chat_input("Ra lệnh..."):
        st.session_state.messages.append({"role": "user", "content": p})
        st.chat_message("user").markdown(p)
        base = f"{knowledge_context}\n{comp_context}" if 'knowledge_context' in locals() else comp_context
        sys = f"{style}\nBắt đầu bằng: '{prefix}'\nDữ liệu: {base}\nCâu hỏi: {p}"
        with st.chat_message("assistant"):
            if AI_STATUS == "Online 🟢":
                try:
                    res = genai.GenerativeModel(MODEL_NAME).generate_content(sys).text
                    st.markdown(res)
                    st.session_state.messages.append({"role": "assistant", "content": res})
                    
                    # Tải file Word
                    docx_file = create_word_docx(current_role_name, res)
                    st.download_button(
                        label="💾 TẢI BẢN WORD (.docx)",
                        data=docx_file,
                        file_name=f"bien_ban_hop_{datetime.now().strftime('%Y%m%d_%H%M')}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key=f"dl_{len(st.session_state.messages)}"
                    )
                except Exception as e: st.error(str(e))
            else: st.error("AI Offline")

elif menu == "⚔️ Rada Đối Thủ":
    st.title("⚔️ RADA ĐỐI THỦ")
    with st.expander("Thêm Đối Thủ"):
        my_l = get_products_list()
        if my_l:
            c1,c2 = st.columns(2)
            p_me = c1.selectbox("SP Mình", my_l)
            p_shop = c1.text_input("Tên Shop")
            p_link = c2.text_input("Link"); p_pr = c2.number_input("Giá", step=1000)
            if st.button("Lưu"): add_competitor(p_me, p_shop, p_link, p_pr); st.rerun()
        else: st.warning("Kho trống!")
    df = get_competitors_df()
    if not df.empty: st.dataframe(df)

elif menu == "💰 Tính Lãi & Nhập Kho":
    st.title("💰 TÍNH LÃI & NHẬP KHO")
    tab1, tab2 = st.tabs(["Thêm Lẻ", "Nhập Excel (Full Data)"])
    with tab1:
        c1,c2,c3=st.columns(3)
        with c1: ten=st.text_input("Tên SP"); von=st.number_input("Giá Vốn", step=1000)
        with c2: ban=st.number_input("Giá Bán", step=1000); hop=st.number_input("Phí gói", 2000)
        with c3: 
            daily=st.number_input("Bán/ngày", 1.0)
            l=st.number_input("Ship (Ngày)", min_value=1, value=5) 
            s=st.number_input("Safe", 5)
        f=st.slider("Phí sàn %",0,30,16)
        if st.button("Tính & Lưu Kho"):
            lai=ban*(1-f/100)-von-hop
            add_product_full(ten, von, ban, 0, daily, l, s) 
            st.metric("Lãi", f"{lai:,.0f}")
            if lai>0: st.success("Đã lưu vào kho!")
    with tab2:
        output = io.BytesIO()
        sample_data = {'Tên sản phẩm': ['Robot T20'], 'Giá vốn': [8000000], 'Giá bán': [12000000], 'Tồn kho': [10], 'Ship (Ngày)': [5], 'Bán/Ngày': [2], 'Tồn An Toàn': [5]}
        df_sample = pd.DataFrame(sample_data)
        with pd.ExcelWriter(output, engine='openpyxl') as writer: df_sample.to_excel(writer, index=False, sheet_name='NhapKho')
        xlsx_sample = output.getvalue()
        
        col_down, col_up = st.columns([1, 2])
        with col_down: st.download_button("📥 Tải File Mẫu (.xlsx)", xlsx_sample, "mau_nhap_kho.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        with col_up:
            f_excel = st.file_uploader("Upload File (.xlsx / .csv)")
            if f_excel:
                if st.button("🚀 Xử Lý Nhập Kho"):
                    try:
                        if f_excel.name.endswith('csv'): df_in = pd.read_csv(f_excel)
                        else: df_in = pd.read_excel(f_excel)
                        count = 0
                        for _, row in df_in.iterrows():
                            try:
                                n = row.get('Tên sản phẩm', row.iloc[0])
                                c = float(row.get('Giá vốn', row.iloc[1]))
                                p = float(row.get('Giá bán', row.iloc[2]))
                                stk = int(row.get('Tồn kho', row.iloc[3])) if len(row) > 3 else 0
                                ship = int(row.get('Ship (Ngày)', row.iloc[4])) if len(row) > 4 else 5
                                dly = float(row.get('Bán/Ngày', row.iloc[5])) if len(row) > 5 else 1.0
                                sfe = int(row.get('Tồn An Toàn', row.iloc[6])) if len(row) > 6 else 5
                                add_product_full(n, c, p, stk, dly, ship, sfe)
                                count += 1
                            except: pass
                        st.success(f"✅ Đã nhập {count} sản phẩm!")
                    except Exception as e: st.error(f"Lỗi: {e}")

elif menu == "📦 Kho Hàng & Backup":
    st.title("📦 QUẢN LÝ KHO & BACKUP")
    df = get_products_df()
    if not df.empty:
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer: df.to_excel(writer, index=False, sheet_name='Backup')
        xlsx_backup = output.getvalue()
        st.download_button("💾 SAO LƯU DỮ LIỆU (.xlsx)", xlsx_backup, "kho_hang_backup.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")
        st.markdown("---")
        st.dataframe(df, use_container_width=True)
    else: st.warning("Kho đang trống.")

